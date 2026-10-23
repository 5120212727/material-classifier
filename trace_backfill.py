# -*- coding: utf-8 -*-
"""将纯精确匹配/KNN结果及历史库溯源信息回填到清单。"""
from __future__ import annotations

import re
import sys
import time
import unicodedata
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.decide.knn import KnnVoteClassifier
from src.ingest.reader import LayoutStore, detect_layout
from src.pipeline import Pipeline, load_config
from src.schema import Item

HEADER_ROW = 2
HEADERS = [
    "是否匹配", "预测大类", "预测中类", "预测小类", "置信度", "置信档",
    "匹配路径", "历史库匹配名称", "历史库匹配规格", "历史库行号",
    "名称对应", "规格对应", "匹配成功的参数", "参数差异/未匹配",
    "是否完全对应", "对应说明",
]
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREEN = Font(color="006100")
RED = Font(color="9C0006", bold=True)

# 「详见技术协议2」「详见技术协议及设备图纸」等变体的兜底识别
_PH_PREFIX = re.compile(
    r"^(详见|参见|见)(技术)?(协议|规格书|规格|图纸|附图|附件|清单|设备表|招标文件|招标|图)"
)


def is_placeholder_spec(s: str, nz) -> bool:
    """规格是否为无参数信息的占位话术（归一化器词表 + 前缀兜底）。"""
    if not s:
        return True
    if nz is not None and nz.is_placeholder(s):
        return True
    return bool(_PH_PREFIX.match(re.sub(r"[\s\u3000]+", "", s).lower()))


def text(value) -> str:
    return "" if pd.isna(value) else str(value).strip()


def parameter_map(spec: str) -> dict[str, list[str]]:
    """提取可解释的工程参数，保留同一参数的多个值。"""
    value = spec or ""
    patterns = [
        ("BW", r"\bBW\s*[=:＝]?\s*(\d+(?:\.\d+)?)\s*(mm)?"),
        ("DN", r"\bDN\s*([0-9]{2,4})"),
        ("PN", r"\bPN\s*([0-9]+(?:\.[0-9]+)?)"),
        ("MPa", r"([0-9]+(?:\.[0-9]+)?)\s*MPa"),
        ("kW", r"(?:\bN\s*[=:＝]?\s*)?([0-9]+(?:\.[0-9]+)?)\s*kW"),
        ("V", r"(?:\bU\s*[=:＝]?\s*)?([0-9]+(?:\.[0-9]+)?)\s*V\b"),
        ("m3/h", r"([0-9]+(?:\.[0-9]+)?)\s*m3\s*/?\s*h"),
        ("m3", r"([0-9]+(?:\.[0-9]+)?)\s*m3(?!\s*/?\s*h)\b"),
        ("mm", r"([0-9]+(?:\.[0-9]+)?)\s*mm\b"),
        ("t", r"([0-9]+(?:\.[0-9]+)?)\s*(?:t\b|吨)"),
        ("进水温度", r"进水温度\s*[:：=＝]?\s*[~～]?\s*([0-9]+(?:\.[0-9]+)?)\s*(?:℃|°C)"),
        ("回水温度", r"回水温度\s*[:：=＝]?\s*[~～]?\s*([0-9]+(?:\.[0-9]+)?)\s*(?:℃|°C)"),
        ("出水温度", r"出水温度\s*[:：=＝]?\s*[~～]?\s*([0-9]+(?:\.[0-9]+)?)\s*(?:℃|°C)"),
    ]
    out: dict[str, list[str]] = {}
    for key, pattern in patterns:
        for match in re.finditer(pattern, value, flags=re.IGNORECASE):
            number = match.group(1)
            unit = "" if key == "BW" else {
                "MPa": "MPa", "kW": "kW", "V": "V", "m3/h": "m3/h",
                "m3": "m3", "mm": "mm", "t": "t",
                "进水温度": "℃", "回水温度": "℃", "出水温度": "℃", "温度": "℃",
            }.get(key, "")
            rendered = f"{number}{unit}"
            if rendered not in out.setdefault(key, []):
                out[key].append(rendered)
    if "BW" in out and "mm" in out:
        out["mm"] = [value for value in out["mm"] if value != f"{out['BW'][0]}mm"]
        if not out["mm"]:
            del out["mm"]
    return out


def format_params(params: dict[str, list[str]]) -> str:
    return "；".join(f"{key}={','.join(values)}" for key, values in params.items())


def compare_params(source_spec: str, history_spec: str) -> tuple[str, str, float]:
    source = parameter_map(source_spec)
    history = parameter_map(history_spec)
    matched: list[str] = []
    differences: list[str] = []
    total_source = sum(len(values) for values in source.values())
    total_matched = 0
    for key in source:
        source_values = source[key]
        history_values = history.get(key, [])
        common = [v for v in source_values if v in history_values]
        total_matched += len(common)
        if common:
            matched.append(f"{key}={','.join(common)}")
        if not common:
            if history_values:
                differences.append(
                    f"{key}：清单 {','.join(source_values)}；历史 {','.join(history_values)}"
                )
            else:
                differences.append(f"{key}：清单 {','.join(source_values)}；历史未找到")
    for key, values in history.items():
        if key not in source:
            differences.append(f"历史多出 {key}={','.join(values)}")
    ratio = total_matched / total_source if total_source else 0.0
    return "；".join(matched) or "无可抽取的共同参数", "；".join(differences) or "无参数差异", ratio


def unique_join(values: list[str], limit: int = 5) -> str:
    unique = []
    for value in values:
        if value and value not in unique:
            unique.append(value)
    if len(unique) > limit:
        return "；".join(unique[:limit]) + f"；等{len(unique)}条"
    return "；".join(unique)


class HistoryIndex:
    """历史库一次性预建名称/规格/分类索引，替代逐条全表布尔掩码扫描。

    - by_name_base：设备主体名称（去位号）-> 历史行位置数组
    - by_candidate：(归一化名称, 归一化规格, 大/中/小类) 组合键 -> 历史行位置数组
    - name_counts / class_counts：同名、同名同类条数统计
    所有查询均为 O(1) 字典命中，位置数组升序，与原布尔掩码结果逐行一致。
    """

    def __init__(self, history: pd.DataFrame, columns: list[str]):
        self.df = history
        self.columns = columns
        self.empty = history.iloc[0:0]
        self.nz = None  # Normalizer，build_history 注入后启用占位规格识别

        # 保持 object dtype（Python str），避免定宽 unicode 数组的内存膨胀
        self._name_base = history["_name_base"].to_numpy()
        self._spec_norm = history["_spec_norm"].to_numpy()
        self._big = history["_big"].to_numpy()
        self._mid = history["_mid"].to_numpy()
        self._small = history["_small"].to_numpy()
        self._name_norm = history["_name_norm"].to_numpy()

        self.by_name_base = self._bin(self._name_base)
        self.name_counts = Counter(self._name_norm)
        self.class_counts = Counter(
            zip(self._name_norm, self._big, self._mid, self._small)
        )

        sep = "\x01"
        # 用 pandas 拼接（object 存储），避免 np.char 定宽 unicode 的内存膨胀
        compound = (
            history["_name_norm"]
            + sep
            + history["_spec_norm"]
            + sep
            + history["_big"]
            + sep
            + history["_mid"]
            + sep
            + history["_small"]
        ).to_numpy()
        self.by_candidate = self._bin(compound)

    @staticmethod
    def _bin(values: np.ndarray) -> dict[str, np.ndarray]:
        unique, inverse = np.unique(values, return_inverse=True)
        order = np.argsort(inverse, kind="stable")
        counts = np.bincount(inverse, minlength=len(unique))
        starts = np.empty(len(unique), dtype=np.int64)
        starts[0] = 0
        starts[1:] = counts[:-1].cumsum()
        return {
            str(unique[i]): order[starts[i] : starts[i] + counts[i]]
            for i in range(len(unique))
        }

    def rows(self, positions) -> pd.DataFrame:
        if positions is None or len(positions) == 0:
            return self.empty
        return self.df.loc[positions]

    def name_base_rows(self, base: str) -> pd.DataFrame:
        return self.rows(self.by_name_base.get(base))

    def candidate_rows(self, candidate) -> pd.DataFrame:
        key = "\x01".join(
            (candidate.src_name, candidate.src_spec, candidate.b, candidate.m, candidate.s)
        )
        return self.rows(self.by_candidate.get(key))

    def same_name_count(self, name_norm: str) -> int:
        return self.name_counts.get(name_norm, 0)

    def same_class_count(self, name_norm: str, b: str, m: str, s: str) -> int:
        return self.class_counts.get((name_norm, b, m, s), 0)


def build_history(path: Path, normalizer) -> HistoryIndex:
    history = pd.read_excel(path)
    columns = list(history.columns)
    history["_name_norm"] = history[columns[0]].map(text).map(normalizer.name)
    history["_name_base"] = history[columns[0]].map(text).map(canonical_name)
    history["_spec_norm"] = history[columns[1]].map(text).map(normalizer.spec)
    history["_big"] = history[columns[2]].map(text)
    history["_mid"] = history[columns[3]].map(text)
    history["_small"] = history[columns[4]].map(text)
    history["_excel_row"] = history.index + 2
    hi = HistoryIndex(history, columns)
    hi.nz = normalizer
    return hi


def canonical_name(value: str) -> str:
    """比较设备主体名称时忽略名称前后 KJ-401/TK301 这类设备位号。"""
    normalized = unicodedata.normalize("NFKC", value or "").strip().lower()
    normalized = re.sub(
        r"^[a-z]{1,3}\s*-?\s*\d{1,4}(?=[\u4e00-\u9fff])",
        "",
        normalized,
    )
    normalized = re.sub(
        r"(?<=[\u4e00-\u9fff])\s*[a-z]{1,3}\s*-?\s*\d{1,4}$",
        "",
        normalized,
    )
    return re.sub(r"[\s，,。;；、·.()（）\[\]【】:：/\\\-—_]+", "", normalized)


def trace_one(item, prediction, exact, knn, hi: HistoryIndex) -> dict:
    columns = hi.columns
    predicted = (prediction.b, prediction.m, prediction.s)
    source_spec_norm = item.spec_norm
    source_name_base = canonical_name(item.name)
    same_base = hi.name_base_rows(source_name_base)
    same_base_label = same_base[
        (same_base["_big"] == prediction.b)
        & (same_base["_mid"] == prediction.m)
        & (same_base["_small"] == prediction.s)
    ]
    exact_spec = same_base[same_base["_spec_norm"] == source_spec_norm]
    exact_same_label = exact_spec[
        (exact_spec["_big"] == prediction.b)
        & (exact_spec["_mid"] == prediction.m)
        & (exact_spec["_small"] == prediction.s)
    ]
    label_conflict = not exact_spec.empty and (
        (exact_spec["_big"] != prediction.b)
        | (exact_spec["_mid"] != prediction.m)
        | (exact_spec["_small"] != prediction.s)
    ).any()

    if not exact_same_label.empty:
        matched_rows = exact_same_label
        full_spec = True
    elif not exact_spec.empty:
        matched_rows = exact_spec
        full_spec = True
    elif not same_base_label.empty:
        ranked_sources = []
        for candidate in exact + knn:
            if candidate.key != predicted:
                continue
            rows = hi.candidate_rows(candidate)
            rows = rows[rows["_name_base"] == source_name_base]
            if not rows.empty:
                ranked_sources.append((candidate.score, rows))
        matched_rows = max(ranked_sources, key=lambda pair: pair[0])[1] if ranked_sources else same_base_label.iloc[[0]]
        full_spec = False
    else:
        matching_candidate = next((c for c in prediction.candidates if c.key == predicted), None)
        if matching_candidate is None and prediction.candidates:
            matching_candidate = prediction.candidates[0]
        matched_rows = hi.candidate_rows(matching_candidate) if matching_candidate else hi.empty
        full_spec = bool(
            source_spec_norm
            and not matched_rows.empty
            and (matched_rows["_spec_norm"] == source_spec_norm).any()
        )

    if matched_rows.empty and prediction.candidates:
        matched_rows = hi.candidate_rows(prediction.candidates[0])
        full_spec = bool(
            source_spec_norm
            and not matched_rows.empty
            and (matched_rows["_spec_norm"] == source_spec_norm).any()
        )

    if exact:
        same_name_count = hi.same_name_count(item.name_norm)
        class_count = hi.same_class_count(item.name_norm, prediction.b, prediction.m, prediction.s)
        if full_spec and not label_conflict:
            path = f"同名精确分类({class_count}/{same_name_count}同名记录同类)；规格检索找到全文一致记录"
        elif full_spec and label_conflict:
            path = "同名精确分类；规格完全对应的历史记录存在标签冲突"
        elif not same_base_label.empty:
            path = "同名精确分类；规格按同名历史记录近似匹配"
        else:
            path = "同名精确分类"
    elif full_spec and not label_conflict:
        path = "KNN加权投票；近邻中找到名称主体+规格全文一致记录"
    elif full_spec and label_conflict:
        path = "KNN加权投票；规格完全对应的历史记录存在标签冲突"
    else:
        path = "KNN Top5 加权投票"

    if not matched_rows.empty:
        # 代表行优选：同一物料的历史记录常有多条，若第一条规格是
        # 「详见技术协议」这类占位话术，优先换到有真实规格的行展示
        if hi.nz is not None:
            raw_specs = matched_rows[columns[1]].map(text)
            has_real = raw_specs.map(lambda s: not is_placeholder_spec(s, hi.nz))
            if has_real.any() and not bool(has_real.iloc[0]):
                matched_rows = pd.concat([matched_rows[has_real], matched_rows[~has_real]])

    if matched_rows.empty:
        history_names = "未找到可回溯的历史原始行"
        history_spec = ""
        history_rows = ""
        representative_name = ""
    else:
        history_names = unique_join(matched_rows[columns[0]].map(text).tolist())
        history_spec = text(matched_rows.iloc[0][columns[1]])
        history_rows = unique_join([str(int(x)) for x in matched_rows["_excel_row"].tolist()], limit=8)
        representative_name = text(matched_rows.iloc[0][columns[0]])

    matched_params, differences, ratio = compare_params(item.spec, history_spec)
    if not matched_rows.empty and any(matched_rows["_name_base"] == source_name_base):
        source_names = set(matched_rows[columns[0]].map(text))
        name_status = "原始名称完全一致" if item.name in source_names else "设备主体名称一致，设备位号/格式不同"
    else:
        name_status = "仅近邻相似"

    if full_spec:
        spec_status = "规格归一化后完全一致"
    elif not matched_rows.empty and ratio > 0:
        spec_status = f"部分对应（参数匹配度 {ratio:.0%}）"
    elif not matched_rows.empty:
        spec_status = "未完全对应"
    else:
        spec_status = "未找到历史规格"

    raw_name_match = not matched_rows.empty and item.name in set(matched_rows[columns[0]].map(text))
    if raw_name_match and full_spec and not label_conflict:
        complete = "是"
        explanation = "清单名称和规格均与历史记录一致"
    elif full_spec and not label_conflict:
        complete = "否（设备位号不同）"
        explanation = "设备主体名称和规格完全一致，但清单与历史记录的设备位号不同"
    elif full_spec and label_conflict:
        complete = "否（历史标签冲突）"
        explanation = "名称主体和规格完全对应，但历史库存在不同分类标签"
    elif not same_base_label.empty:
        complete = "否"
        explanation = "设备主体名称对应，但规格没有完全对应"
    elif matched_rows.empty:
        complete = "否"
        explanation = "只有KNN分类近邻，未找到可回溯的完整历史记录"
    else:
        complete = "否"
        explanation = "只有KNN近邻，名称或规格没有完全对应"

    return {
        "物料名称": item.name,
        "清单规格": item.spec,
        "源行号": item.idx,
        "是否匹配": "是" if prediction.band == "高" else ("待抽查" if prediction.band == "中" else "待人工确认"),
        "预测大类": prediction.b,
        "预测中类": prediction.m,
        "预测小类": prediction.s,
        "置信度": round(prediction.confidence, 3),
        "置信档": prediction.band,
        "匹配路径": path,
        "历史库匹配名称": history_names,
        "历史库匹配规格": history_spec,
        "历史库行号": history_rows,
        "名称对应": name_status,
        "规格对应": spec_status,
        "匹配成功的参数": matched_params,
        "参数差异/未匹配": differences,
        "是否完全对应": complete,
        "对应说明": explanation,
    }


def resolve_header_row(wb, sheet_name: str, layout_dir: Path) -> int:
    """优先读布局配置，其次自动探测；都失败时退回默认第 2 行。"""
    store = LayoutStore(layout_dir)
    lay = store.load(sheet_name) or detect_layout(wb[sheet_name], sheet_name)
    return lay.header_row if lay is not None else HEADER_ROW


def main() -> int:
    if len(sys.argv) not in (4, 5):
        print("用法: python trace_backfill.py <源清单.xlsx> <knn结果.xlsx> <输出.xlsx> [页签名，默认 炼铁]")
        return 2
    source_path, result_path, output_path = map(Path, sys.argv[1:4])
    sheet_name = sys.argv[4] if len(sys.argv) == 5 else "炼铁"
    base = Path(__file__).resolve().parent
    t_start = time.perf_counter()
    cfg = load_config(str(base / "config" / "default.yaml"))
    cfg["decide"]["strategy"] = "knn_vote"
    pipe = Pipeline(cfg).build()
    t_pipe = time.perf_counter()

    result = pd.read_excel(result_path)
    items = []
    for _, row in result.iterrows():
        item = Item(
            idx=int(row["源行号"]),
            name=text(row["物料名称"]),
            spec=text(row["型号、规格及技术性能"]),
            sheet=text(row["页签"]),
        )
        item.name_norm = pipe.nz.name(item.name)
        item.spec_norm = pipe.nz.spec(item.spec)
        items.append(item)

    t_search = time.perf_counter()
    exact, knn = pipe.kb.search(items, topk=10)
    classifier = KnnVoteClassifier(pipe.kb, cfg["decide"].get("knn_vote", {}))
    predictions = classifier.classify(items, exact, knn)
    t_classify = time.perf_counter()

    hi = build_history(Path(cfg["knowledge_base"]), pipe.nz)
    t_index = time.perf_counter()

    traces = [trace_one(item, pred, exact[i], knn[i], hi) for i, (item, pred) in enumerate(zip(items, predictions))]
    t_trace = time.perf_counter()
    trace_df = pd.DataFrame(traces)

    wb = load_workbook(source_path)
    ws = wb[sheet_name]
    header_row = resolve_header_row(wb, sheet_name, Path(cfg.get("layout_dir", "config/layouts")))
    base_column = ws.max_column
    for offset, header in enumerate(HEADERS, 1):
        cell = ws.cell(header_row, base_column + offset, header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    widths = [14, 16, 18, 18, 10, 10, 24, 28, 48, 18, 24, 24, 34, 42, 18, 32]
    for offset, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(base_column + offset)].width = width

    for result_row, trace in zip(result.itertuples(index=False), trace_df.to_dict("records")):
        source_row = int(getattr(result_row, "源行号"))
        band = trace["置信档"]
        for offset, header in enumerate(HEADERS, 1):
            value = trace[header]
            cell = ws.cell(source_row, base_column + offset, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if band == "低":
                cell.fill = YELLOW
            if header == "是否匹配":
                cell.font = RED if band == "低" else GREEN

    wb.save(output_path)
    trace_df.to_csv(output_path.with_suffix(".csv"), index=False, encoding="utf-8-sig")
    t_done = time.perf_counter()
    print(f"· 纯规则分类: {len(predictions)} 条")
    print("· 置信档:", dict(trace_df["置信档"].value_counts()))
    print("· 完全对应:", dict(trace_df["是否完全对应"].value_counts()))
    print("· 匹配路径:", dict(trace_df["匹配路径"].value_counts()))
    print(
        f"· 耗时: 初始化 {t_pipe - t_start:.1f}s | 读取+归一化 {t_search - t_pipe:.1f}s | "
        f"KNN检索+分类 {t_classify - t_search:.1f}s | 历史库索引构建 {t_index - t_classify:.1f}s | "
        f"追溯回填 {t_trace - t_index:.1f}s | Excel写入 {t_done - t_trace:.1f}s | 合计 {t_done - t_start:.1f}s"
    )
    print(f"· 输出: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
