"""输入适配层：把任意格式的报价清单 Excel 解析成统一 Item 列表。

核心思路：**布局用配置描述，不用代码写死。**
- 已知专业表 → config/layouts/*.yaml
- 未知表 → 自动探测表头，探测结果可回写成 yaml 复用
- 只有在自动探测失败时，才需要人工介入（这一处可以接 LLM，见 decide 层之外的 hook）
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
import yaml

from ..schema import Item, TableLayout

# 名称列 / 规格列的候选关键词（按优先级）
NAME_KEYS = ["工程或费用名称", "费用名称", "材料名称", "设备名称", "名称", "物料名称"]
SPEC_KEYS = ["型号、规格及技术性能", "型号规格及技术性能", "型号、规格", "规格型号", "规格", "型号"]


def _norm_header(s) -> str:
    if s is None:
        return ""
    s = str(s).replace("\n", "").replace(" ", "").replace("\u3000", "")
    return s.strip()


def _match(cell: str, keys: list[str]) -> int:
    """返回命中关键词的优先级序号，越小越优先；未命中返回 -1。"""
    for i, k in enumerate(keys):
        if k and k in cell:
            return i
    return -1


class LayoutStore:
    """布局配置仓库。"""

    def __init__(self, layout_dir: str | Path = "config/layouts"):
        self.dir = Path(layout_dir)
        self.dir.mkdir(parents=True, exist_ok=True)

    def load(self, sheet: str) -> TableLayout | None:
        for suf in (".yaml", ".yml"):
            p = self.dir / f"{sheet}{suf}"
            if p.exists():
                d = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
                return TableLayout(
                    sheet=sheet,
                    header_row=int(d.get("header_row", 2)),
                    data_start_row=int(d.get("data_start_row", 4)),
                    name_col=int(d["name_col"]),
                    spec_col=int(d["spec_col"]),
                    skip_keywords=d.get("skip_keywords", ["小计", "合计", "总计", "合价"]),
                    layout_name=str(p.stem),
                )
        return None

    def save(self, layout: TableLayout) -> Path:
        p = self.dir / f"{layout.sheet}.yaml"
        p.write_text(
            yaml.safe_dump(
                {
                    "header_row": layout.header_row,
                    "data_start_row": layout.data_start_row,
                    "name_col": layout.name_col,
                    "spec_col": layout.spec_col,
                    "skip_keywords": layout.skip_keywords,
                },
                allow_unicode=True,
                sort_keys=False,
            ),
            encoding="utf-8",
        )
        return p


def detect_layout(ws, sheet: str, probe_rows: int = 8) -> TableLayout | None:
    """自动探测表头：在前若干行里找同时含名称列与规格列的那一行。"""
    best = None
    for r in range(1, min(probe_rows, ws.max_row) + 1):
        cells = [_norm_header(ws.cell(row=r, column=c).value) for c in range(1, min(ws.max_column, 40) + 1)]
        name_hits = [(_match(c, NAME_KEYS), i + 1) for i, c in enumerate(cells) if _match(c, NAME_KEYS) >= 0]
        spec_hits = [(_match(c, SPEC_KEYS), i + 1) for i, c in enumerate(cells) if _match(c, SPEC_KEYS) >= 0]
        if name_hits and spec_hits:
            name_hits.sort()
            # 规格列排除「型号、规格及技术性能-调整」这类派生列，取最左
            spec_hits.sort(key=lambda t: (t[1], t[0]))
            cand = (r, name_hits[0][1], spec_hits[0][1])
            if best is None:
                best = cand
            # 优先取带「工程或费用名称」的精确匹配
            if name_hits[0][0] == 0:
                best = cand
                break
    if best is None:
        return None
    hr, nc, sc = best
    # 数据起始行：从表头之后找到第一行 name 与 spec 同时非空
    dsr = hr + 1
    for r in range(hr + 1, min(hr + 12, ws.max_row) + 1):
        v1 = ws.cell(row=r, column=nc).value
        v2 = ws.cell(row=r, column=sc).value
        if v1 not in (None, "") and v2 not in (None, ""):
            dsr = r
            break
    return TableLayout(sheet=sheet, header_row=hr, data_start_row=dsr, name_col=nc, spec_col=sc)


def read_workbook(
    path: str | Path,
    sheets: list[str] | None = None,
    layout_dir: str | Path = "config/layouts",
    save_detected: bool = False,
    exclude_sheets: list[str] | None = None,
) -> tuple[list[Item], list[TableLayout]]:
    """读取清单，返回 (物料列表, 使用的布局列表)。"""
    store = LayoutStore(layout_dir)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    items: list[Item] = []
    layouts: list[TableLayout] = []
    ex = set(exclude_sheets or [])
    try:
        targets = sheets or wb.sheetnames
        for sn in targets:
            if sn not in wb.sheetnames or sn in ex:
                continue
            ws = wb[sn]
            lay = store.load(sn) or detect_layout(ws, sn)
            if lay is None:
                continue
            if save_detected and store.load(sn) is None:
                store.save(lay)
            layouts.append(lay)
            for r in range(lay.data_start_row, ws.max_row + 1):
                name = ws.cell(row=r, column=lay.name_col).value
                spec = ws.cell(row=r, column=lay.spec_col).value
                name = "" if name is None else str(name).strip()
                spec = "" if spec is None else str(spec).strip()
                if not name or not spec:
                    continue
                # 剔除分组行 / 小计行
                if any(k in name for k in lay.skip_keywords) and len(spec) < 4:
                    continue
                items.append(Item(idx=r, name=name, spec=spec, sheet=sn))
    finally:
        wb.close()
    return items, layouts


def inspect(path: str | Path, sheet: str | None = None) -> None:
    """调试用：打印每个 sheet 的探测结果。"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for sn in wb.sheetnames:
        if sheet and sn != sheet:
            continue
        ws = wb[sn]
        lay = detect_layout(ws, sn)
        if lay:
            print(f"[{sn}] header_row={lay.header_row} data_start={lay.data_start_row} "
                  f"name_col={lay.name_col} spec_col={lay.spec_col}")
        else:
            print(f"[{sn}] 无法自动探测（需要人工配置 layout）")
    wb.close()
