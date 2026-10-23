#!/usr/bin/env python3
"""评测 runner：在同一套留出集上对比不同策略 / 不同模型。

关键设计：**留出集在切分前先按 (名称,规格) 去重**，
所以测试样本在知识库里绝对不存在完全相同的记录 —— 测的是「对没见过的物料」的能力。

用法：
  python -m eval.runner                        # 只跑规则基线
  python -m eval.runner -s knn_vote llm_rerank -m deepseek-chat
  python -m eval.runner -s llm_direct llm_rerank -m qwen-max --holdout 0.15
"""
from __future__ import annotations

import argparse
import json
import random
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.decide.knn import KnnVoteClassifier  # noqa: E402
from src.decide.llm_direct import LLMDirectClassifier  # noqa: E402
from src.decide.llm_rerank import LLMRerankClassifier  # noqa: E402
from src.llm.client import load_client  # noqa: E402
from src.normalize.rules import Normalizer  # noqa: E402
from src.pipeline import load_config  # noqa: E402
from src.retrieve.index import KnowledgeBase, _is_dirty  # noqa: E402
from src.schema import Item  # noqa: E402

FACTORY = {"knn_vote": KnnVoteClassifier, "llm_rerank": LLMRerankClassifier, "llm_direct": LLMDirectClassifier}


def build_holdout(hist_path: str, nz: Normalizer, ratio: float, seed: int):
    """返回 (测试样本列表, 测试集 key 集合, 全量统计)。"""
    wb = openpyxl.load_workbook(hist_path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    raw = []
    for r in ws.iter_rows(min_row=2, max_col=5, values_only=True):
        if not r:
            continue
        r = tuple(r) + (None,) * 5
        nm, sp, b, m, s = [(str(x).strip() if x is not None else "") for x in r[:5]]
        if b and m and s:
            raw.append((nm, sp, b, m, s))
    wb.close()

    grp: dict[tuple[str, str], list] = defaultdict(list)
    for nm, sp, b, m, s in raw:
        grp[(nz.name(nm), nz.spec(sp))].append((nm, sp, b, m, s))

    recs = []
    for (k1, k2), rows in grp.items():
        c = Counter((x[2], x[3], x[4]) for x in rows)
        (b, m, s), _ = c.most_common(1)[0]
        if _is_dirty(b, m, s):
            continue
        nm, sp = rows[0][0], rows[0][1]
        recs.append({"nk": k1, "sk": k2, "name": nm, "spec": sp, "b": b, "m": m, "s": s})

    rng = random.Random(seed)
    idx = list(range(len(recs)))
    rng.shuffle(idx)
    cut = int(len(recs) * (1 - ratio))
    test = [recs[i] for i in idx[cut:]]
    test_keys = {(r["nk"], r["sk"]) for r in test}
    return test, test_keys, {"total": len(recs), "test": len(test), "train": cut}


def to_items(test: list[dict], nz: Normalizer) -> list[Item]:
    out = []
    for i, r in enumerate(test):
        out.append(Item(idx=i + 1, name=r["name"], spec=r["spec"],
                        name_norm=r["nk"], spec_norm=r["sk"], sheet="holdout"))
    return out


def evaluate(preds, test) -> dict:
    n = len(test)
    okb = sum(1 for p, r in zip(preds, test) if p.b == r["b"])
    okm = sum(1 for p, r in zip(preds, test) if p.m == r["m"])
    oks = sum(1 for p, r in zip(preds, test) if p.s == r["s"])
    oka = sum(1 for p, r in zip(preds, test) if p.b == r["b"] and p.m == r["m"] and p.s == r["s"])
    cov = sum(1 for p in preds if p.b)
    return {"n": n, "cov": cov / n, "b": okb / n, "m": okm / n, "s": oks / n, "all": oka / n}


def confidence_curve(preds, test, ths=(0.3, 0.4, 0.5, 0.6, 0.7, 0.8)) -> list[dict]:
    out = []
    for th in ths:
        keep = [(p, r) for p, r in zip(preds, test) if p.b and p.confidence >= th]
        if not keep:
            continue
        m = len(keep)
        out.append({
            "th": th, "cov": m / len(test),
            "b": sum(1 for p, r in keep if p.b == r["b"]) / m,
            "m": sum(1 for p, r in keep if p.m == r["m"]) / m,
            "s": sum(1 for p, r in keep if p.s == r["s"]) / m,
        })
    return out


def error_profile(preds, test, topn=8) -> list[tuple]:
    """错在哪些大类上，便于定位问题。"""
    c = Counter()
    for p, r in zip(preds, test):
        if p.b != r["b"]:
            c[(r["b"], p.b or "(空)")] += 1
    return c.most_common(topn)


def run_one(strategy: str, model: str | None, cfg: dict, items, test, test_keys, nz) -> tuple[dict, list]:
    kb = KnowledgeBase(
        cfg["knowledge_base"], nz,
        cache_dir=cfg.get("cache_dir", "cache"),
        name_weight=float(cfg.get("retrieve", {}).get("name_weight", 0.75)),
        spec_weight=float(cfg.get("retrieve", {}).get("spec_weight", 0.25)),
        min_df=int(cfg.get("retrieve", {}).get("min_df", 2)),
        exclude_keys=test_keys,
    ).build()

    params = dict(cfg.get("decide", {}).get(strategy, {}))
    params.setdefault("topk", cfg.get("retrieve", {}).get("topk", 10))
    if strategy.startswith("llm_"):
        if not model:
            raise SystemExit(f"策略 {strategy} 需要 --model")
        clf = FACTORY[strategy](kb, params, client=load_client(model, cfg.get("models_path")))
    else:
        clf = FACTORY[strategy](kb, params)

    t0 = time.time()
    exact, knn = kb.search(items, topk=int(cfg.get("retrieve", {}).get("topk", 10)))
    preds = clf.classify(items, exact, knn)
    el = time.time() - t0
    res = evaluate(preds, test)
    res.update({"strategy": strategy, "model": model or "-", "elapsed_s": round(el, 2),
                "per_item_ms": round(el / len(items) * 1000, 1),
                "curve": confidence_curve(preds, test)})
    return res, preds


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--config", "-c", default="config/default.yaml")
    ap.add_argument("--strategies", "-s", nargs="+", default=["knn_vote"],
                    choices=list(FACTORY))
    ap.add_argument("--model", "-m", default=None)
    ap.add_argument("--holdout", type=float, default=0.10)
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--json", default="out/eval_report.json")
    args = ap.parse_args()

    cfg = load_config(args.config)
    nz = Normalizer(strip_tags=True, drop_placeholders=True)

    print(f"[1/3] 构建留出集（去重后随机留出 {args.holdout:.0%}）...")
    test, test_keys, meta = build_holdout(cfg["knowledge_base"], nz, args.holdout, args.seed)
    print(f"      唯一记录 {meta['total']:,} → 训练 {meta['train']:,} / 测试 {meta['test']:,}")
    items = to_items(test, nz)

    print(f"[2/3] 逐个策略评测：{', '.join(args.strategies)}")
    results = []
    all_preds = {}
    for s in args.strategies:
        r, preds = run_one(s, args.model, cfg, items, test, test_keys, nz)
        all_preds[s] = preds
        results.append(r)
        print(f"      {s:12s} 覆盖 {r['cov']*100:5.1f}% | 大类 {r['b']*100:5.1f}% | "
              f"中类 {r['m']*100:5.1f}% | 小类 {r['s']*100:5.1f}% | 全对 {r['all']*100:5.1f}% | "
              f"{r['per_item_ms']:.1f} ms/条")

    print("[3/3] 置信度曲线")
    for r in results:
        print(f"      [{r['strategy']}]")
        for c in r["curve"]:
            print(f"        阈值 {c['th']:.1f}  覆盖 {c['cov']*100:5.1f}%  "
                  f"大类 {c['b']*100:5.1f}%  中类 {c['m']*100:5.1f}%  小类 {c['s']*100:5.1f}%")
        errs = error_profile(all_preds[r["strategy"]], test)
        if errs:
            print("        主要误判：" + "；".join(f"真{tr}→判{pr}({n})" for (tr, pr), n in errs[:5]))

    out = Path(args.json)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({"meta": meta, "results": results}, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"\n报告已写入 {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
