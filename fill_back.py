# -*- coding: utf-8 -*-
"""把分类结果回填到清单副本的专业表上（独立工具，不改 agent 本体）。

用法: python fill_back.py <清单xlsx> <结果xlsx> <输出xlsx> <页签名1,页签名2,...>
"""
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_ROW = 2
NEW_HEADERS = ["是否匹配", "预测大类", "预测中类", "预测小类", "置信度", "置信档"]
YELLOW = PatternFill("solid", fgColor="FFF2CC")
RED = Font(color="9C0006", bold=True)
GREEN = Font(color="006100")


def main() -> None:
    src, res_file, dst = sys.argv[1], sys.argv[2], sys.argv[3]
    sheets = sys.argv[4].split(",")
    res = pd.read_excel(res_file)
    cols = res.columns
    by_sheet = {s: res[res[cols[1]] == s] for s in sheets}

    wb = load_workbook(src)
    for sheet, sub in by_sheet.items():
        ws = wb[sheet]
        base = ws.max_column
        for i, h in enumerate(NEW_HEADERS):
            ws.cell(row=HEADER_ROW, column=base + 1 + i, value=h).font = Font(bold=True)
        for _, r in sub.iterrows():
            row = int(r[cols[2]])
            conf = float(r[cols[8]])
            band = str(r[cols[9]])
            reason = str(r[cols[11]])
            if band == "高":
                flag = "是"
            elif band == "中":
                flag = "待抽查"
            else:
                flag = "知识库无此物料" if "拒识" in reason else "待人工确认"
            vals = [flag, r[cols[5]], r[cols[6]], r[cols[7]], conf, band]
            for i, v in enumerate(vals):
                if pd.isna(v):
                    v = ""
                cc = ws.cell(row=row, column=base + 1 + i, value=v)
                if band == "低":
                    cc.fill = YELLOW
                if i == 0:
                    cc.font = RED if band == "低" else GREEN
            ws.column_dimensions[get_column_letter(base + 1)].width = 16
        print(f"[{sheet}] 回填 {len(sub)} 行，新增列从第 {base + 1} 列开始")
    wb.save(dst)
    print("saved:", dst)


if __name__ == "__main__":
    main()
