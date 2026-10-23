# -*- coding: utf-8 -*-
"""一键物料分类：自动探测布局 -> (清理宏警告残留) -> LLM 重载 -> 结果回填文档。

用法:
  python auto_classify.py <报价清单.xlsx> [--sheets 炼铁,给排水] [--model qwen3.8-27b]
                          [--max-llm-items 0] [--out out/自定义.xlsx]

依赖环境变量 AIROUTER_API_KEY（未设置时自动降级为免 API 的 knn_vote 策略）。
"""
import argparse
import os
import re
import sys
import threading
import time
import zipfile
from pathlib import Path

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))

import pandas as pd  # noqa: E402
import requests  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from src.ingest.reader import LayoutStore, detect_layout  # noqa: E402
from src.llm import client as llm_client  # noqa: E402
from src.pipeline import Pipeline, load_config  # noqa: E402

NEW_HEADERS = ["是否匹配", "预测大类", "预测中类", "预测小类", "置信度", "置信档"]
YELLOW = PatternFill("solid", fgColor="FFF2CC")
RED = Font(color="9C0006", bold=True)
GREEN = Font(color="006100")


# ---------- 1. 清理宏警告残留 ----------
def sanitize(src: Path, work_dir: Path) -> Path:
    zf = zipfile.ZipFile(src)
    wbxml = zf.read("xl/workbook.xml").decode("utf-8")
    need = any(n.startswith("xl/externalLinks") for n in zf.namelist())
    need = need or "EVALUATE(" in wbxml or "#NAME?" in wbxml
    if not need:
        print("· 文档干净，无需清理")
        return src
    items = {i.filename: zf.read(i.filename) for i in zf.infolist()}
    for n in [n for n in items if n.startswith("xl/externalLinks")]:
        del items[n]
    ct = items["[Content_Types].xml"].decode("utf-8")
    key = "<Override PartName=\"/xl/externalLinks/"
    a = ct.find(key)
    while a >= 0:
        b = ct.find("/>", a)
        ct = ct[:a] + ct[b + 2:]
        a = ct.find(key, a)
    items["[Content_Types].xml"] = ct.encode("utf-8")
    rt = items["xl/_rels/workbook.xml.rels"].decode("utf-8")
    a = rt.find("<Relationship")
    while a >= 0:
        b = rt.find("/>", a)
        if "externalLink" in rt[a:b + 2]:
            rt = rt[:a] + rt[b + 2:]
        else:
            a = rt.find("<Relationship", b + 2)
    items["xl/_rels/workbook.xml.rels"] = rt.encode("utf-8")
    wb = re.sub(r"<externalReferences>.*?</externalReferences>", "", wbxml)

    def keep_dn(m):
        frag = m.group(0)
        return "" if ("#REF!" in frag or "#NAME?" in frag or "EVALUATE(" in frag) else frag

    wb = re.sub(r"<definedName [^>]*>.*?</definedName>", keep_dn, wb)
    items["xl/workbook.xml"] = wb.encode("utf-8")
    dst = work_dir / (src.stem + "_清理版.xlsx")
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as out:
        for n, d in items.items():
            out.writestr(n, d)
    print(f"· 检测到宏警告残留，已清理 -> {dst.name}")
    return dst


# ---------- 2. 布局预览 ----------
def preview_layouts(path: Path, layout_dir: str) -> list[str]:
    wb = load_workbook(path)
    store = LayoutStore(layout_dir)
    ok = []
    for sn in wb.sheetnames:
        lay = store.load(sn) or detect_layout(wb[sn], sn)
        if lay is None:
            print(f"  [跳过] {sn}: 无法自动探测（汇总表或新格式）")
        else:
            print(f"  [处理] {sn}: 表头第{lay.header_row}行, 数据第{lay.data_start_row}行起, 名称列{lay.name_col}, 规格列{lay.spec_col}")
            ok.append(sn)
    wb.close()
    return ok


# ---------- 3. LLM 客户端加料：进度打点 + 扩大连接池 ----------
_counter = {"done": 0, "fail": 0, "lock": threading.Lock()}
_orig_chat = llm_client.LLMClient.chat


def _patched_chat(self, system, user):
    t0 = time.time()
    r = _orig_chat(self, system, user)
    with _counter["lock"]:
        _counter["done"] += 1
        if r is None:
            _counter["fail"] += 1
        n = _counter["done"]
        if n % 10 == 0:
            print(f"  LLM [{n}] ok={_counter['done']-_counter['fail']} fail={_counter['fail']} last={time.time()-t0:.1f}s", flush=True)
    return r


def _patched_init(self, cfg):
    self.cfg = cfg
    self.key = cfg.resolve_key()
    s = requests.Session()
    ad = requests.adapters.HTTPAdapter(pool_connections=64, pool_maxsize=64)
    s.mount("http://", ad)
    s.mount("https://", ad)
    self._session = s


llm_client.LLMClient.chat = _patched_chat
llm_client.LLMClient.__init__ = _patched_init


# ---------- 4. 回填 ----------
def backfill(doc_src: Path, res_path: Path, out_path: Path, sheets: list[str]) -> None:
    res = pd.read_excel(res_path)
    cols = res.columns
    wb = load_workbook(doc_src)
    for sheet in sheets:
        sub = res[res[cols[1]] == sheet]
        if not len(sub):
            continue
        ws = wb[sheet]
        base = ws.max_column
        for i, h in enumerate(NEW_HEADERS):
            ws.cell(row=2, column=base + 1 + i, value=h).font = Font(bold=True)
        for _, r in sub.iterrows():
            row = int(r[cols[2]])
            band = str(r[cols[9]])
            reason = str(r[cols[11]])
            if band == "高":
                flag = "是"
            elif band == "中":
                flag = "待抽查"
            else:
                flag = "知识库无此物料" if "拒识" in reason else "待人工确认"
            vals = [flag, r[cols[5]], r[cols[6]], r[cols[7]], float(r[cols[8]]), band]
            for i, v in enumerate(vals):
                if pd.isna(v):
                    v = ""
                cc = ws.cell(row=row, column=base + 1 + i, value=v)
                if band == "低":
                    cc.fill = YELLOW
                if i == 0:
                    cc.font = RED if band == "低" else GREEN
        ws.column_dimensions[get_column_letter(base + 1)].width = 16
        print(f"  [回填] {sheet}: {len(sub)} 行, 新增列从第 {base + 1} 列起")
    wb.save(out_path)


# ---------- 主流程 ----------
def main() -> int:
    ap = argparse.ArgumentParser(description="一键物料分类（自动探测 -> LLM 重载 -> 回填文档）")
    ap.add_argument("input", help="报价清单 xlsx")
    ap.add_argument("--sheets", default="", help="只处理指定页签,逗号分隔;默认全部可探测页签")
    ap.add_argument("--model", default="qwen3.8-27b")
    ap.add_argument("--max-llm-items", type=int, default=0, help="0=全部走 LLM")
    ap.add_argument("--out", default="")
    args = ap.parse_args()

    src = Path(args.input).resolve()
    work = BASE / "out"
    work.mkdir(exist_ok=True)

    print(f"· 源文件: {src.name}")
    clean = sanitize(src, work)

    cfg = load_config(str(BASE / "config/default.yaml"))
    excl = set((cfg.get("sheets") or {}).get("exclude", []))
    print("· 布局探测:")
    detectable = preview_layouts(clean, str(BASE / "config" / "layouts"))
    want = [s.strip() for s in args.sheets.split(",") if s.strip()]
    if want:
        missing = [s for s in want if s not in detectable]
        if missing:
            print(f"  !! 无法探测: {missing}")
        sheets = [s for s in want if s in detectable]
    else:
        sheets = [s for s in detectable if s not in excl]
    if not sheets:
        print("没有可处理的页签")
        return 1

    has_key = bool(os.environ.get("AIROUTER_API_KEY"))
    strategy = "llm_rerank" if has_key else "knn_vote"
    if not has_key:
        print("· 未设置 AIROUTER_API_KEY, 降级为 knn_vote（免 API）")
    print(f"· 策略: {strategy}" + (f" ({args.model})" if strategy != "knn_vote" else ""))

    cfg["decide"]["strategy"] = strategy
    if strategy != "knn_vote":
        cfg["decide"]["model"] = args.model
        cfg["decide"].setdefault(strategy, {})["max_llm_items"] = args.max_llm_items

    pipe = Pipeline(cfg).build()
    res_path = work / f"{src.stem}_{strategy}_auto.xlsx"
    out_path = Path(args.out) if args.out else work / f"{src.stem}_已标注_{strategy}.xlsx"

    t0 = time.time()
    items, preds, st = pipe.run(str(clean), sheets=sheets, out_path=str(res_path))
    print("· 分类完成: " + f"{st.total} 条, 用时 {st.elapsed_s:.1f}s, 置信档: "
          + " ".join(f"{k} {v}" for k, v in sorted(st.by_band.items())))
    print("· 判定路径: " + " ".join(f"{k} {v}" for k, v in sorted(st.by_strategy.items(), key=lambda x: -x[1])))

    backfill(clean, res_path, out_path, sheets)
    print(f"· 总耗时 {time.time()-t0:.0f}s")
    print(f"· 最终交付: {out_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
